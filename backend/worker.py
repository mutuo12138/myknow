from __future__ import annotations

import html
import ipaddress
import json
import os
import re
import socket
import sqlite3
import time
import uuid
import urllib.request
from pathlib import Path
from urllib.parse import urlparse

from PIL import Image, ImageOps

from ai_adapter import analyze_material, compact_tags, redact_sensitive_text
from main import (
    ASSET_ROOT, STAGING_ROOT, ai_usage_today, category_for, db, finish_ai_batch_item,
    get_processing_settings, init_db, mark_operation_failed, now, queue_job, refresh_ai_batch,
    safe_folder_name, schedule_auto_batch, sha256_bytes, start_ai_batch_item, write_staged_file,
)

MAX_IMAGE_BYTES = 30 * 1024 * 1024
THUMBNAIL_MAX_EDGE = 640
_restart_recovery_done = False


def image_prompt_bundle(result: dict) -> str:
    """Store the bounded reverse-prompt set in the existing description field."""
    positive = str(result.get("prompt_positive") or "").strip()[:900]
    negative = str(result.get("prompt_negative") or "").strip()[:360]
    checklist = str(result.get("verification_checklist") or "").strip()[:360]
    description = str(result.get("visual_description") or "").strip()[:1600]
    if not any((positive, negative, checklist)):
        return description[:2000]
    return "\n".join(part for part in (
        f"Positive Prompt: {positive}" if positive else "",
        f"Negative Prompt: {negative}" if negative else "",
        f"Verification: {checklist}" if checklist else "",
        f"观察摘要：{description}" if description else "",
    ) if part)[:2000]


def is_ai_provider_error(error: Exception) -> bool:
    """AI failures are user-actionable; never hide them behind silent retries."""
    message = str(error).lower()
    return any(marker in message for marker in (
        "ai api", "insufficient_quota", "quota exhausted", "unsupported model",
        "invalid model", "model_not_found", " 401", " 403", " 429",
        "返回了空内容", "没有返回 json", "无法连接",
    ))


def update_fts(conn: sqlite3.Connection, material_id: str) -> None:
    conn.execute("DELETE FROM materials_fts WHERE material_id=?", (material_id,))
    row = conn.execute(
        """SELECT m.title, m.text_content, m.ai_summary, m.tags_json, m.deleted_at,
                  v.visual_description
           FROM materials m LEFT JOIN visual_features v ON v.material_id=m.id WHERE m.id=?""",
        (material_id,),
    ).fetchone()
    if not row or row["deleted_at"] is not None:
        return
    conn.execute("INSERT INTO materials_fts(material_id, title, text_content, ai_summary, visual_description, tags) VALUES (?, ?, ?, ?, ?, ?)",
                 (material_id, row["title"], row["text_content"] or "", row["ai_summary"] or "", row["visual_description"] or "", " ".join(json.loads(row["tags_json"] or "[]"))))


def analyze_visual(path: str) -> tuple[list[dict], dict, float, float]:
    with Image.open(path) as image:
        image = image.convert("RGB")
        image.thumbnail((256, 256))
        pixels = list(image.getdata())
        if not pixels:
            return [], {"r": [], "g": [], "b": [], "bins": 64}, 0.0, 0.0
        mean = sum((red + green + blue) / (3 * 255) for red, green, blue in pixels) / len(pixels)
        buckets: dict[tuple[int, int, int], int] = {}
        sample = pixels[::max(1, len(pixels) // 16000)]
        for red, green, blue in sample:
            key = (red // 32 * 32, green // 32 * 32, blue // 32 * 32)
            buckets[key] = buckets.get(key, 0) + 1
        total = sum(buckets.values()) or 1
        palette = [
            {"rgb": color, "hex": "#%02X%02X%02X" % color, "ratio": round(count / total, 3)}
            for color, count in sorted(buckets.items(), key=lambda item: item[1], reverse=True)[:5]
        ]

        # A photography histogram compares the red, green and blue channels on
        # one shared vertical scale. Keep 64 bins so the chart stays compact
        # in the inspector while preserving the tonal shape users expect.
        channel_bins = {channel: [0] * 64 for channel in ("r", "g", "b")}
        for red, green, blue in sample:
            channel_bins["r"][min(63, red * 64 // 256)] += 1
            channel_bins["g"][min(63, green * 64 // 256)] += 1
            channel_bins["b"][min(63, blue * 64 // 256)] += 1
        peak = max((max(values) for values in channel_bins.values()), default=0) or 1
        histogram = {
            "bins": 64,
            "r": [round(value / peak, 4) for value in channel_bins["r"]],
            "g": [round(value / peak, 4) for value in channel_bins["g"]],
            "b": [round(value / peak, 4) for value in channel_bins["b"]],
        }
        luminance = [sum(channel_bins[channel][index] for channel in ("r", "g", "b")) / (peak * 3) for index in range(64)]
        detail = sum(abs(luminance[index] - luminance[index - 1]) for index in range(1, 64)) / 8
        return palette, histogram, round(mean, 3), round(min(detail, 1), 3)


def generate_thumbnail(material: dict) -> tuple[str, int, int]:
    source_path = Path(material.get("local_path") or "")
    if not source_path.is_file():
        raise FileNotFoundError("生成缩略图所需的本地文件不存在")
    folder = ASSET_ROOT / "derivatives" / "thumbnail" / safe_folder_name(material.get("folder_id") or "inbox")
    final_path = folder / f"{material['id']}.webp"
    staging_dir = STAGING_ROOT / f"thumbnail-{material['id']}-{uuid.uuid4().hex[:8]}"
    staged_path = staging_dir / "thumbnail.webp"
    staging_dir.mkdir(parents=True, exist_ok=False)
    try:
        with Image.open(source_path) as source:
            source.seek(0)
            image = ImageOps.exif_transpose(source)
            if image.mode in {"RGBA", "LA"} or (image.mode == "P" and "transparency" in image.info):
                rgba = image.convert("RGBA")
                image = Image.new("RGB", rgba.size, "white")
                image.paste(rgba, mask=rgba.getchannel("A"))
            else:
                image = image.convert("RGB")
            image.thumbnail((THUMBNAIL_MAX_EDGE, THUMBNAIL_MAX_EDGE), Image.Resampling.LANCZOS)
            width, height = image.size
            if width < 1 or height < 1:
                raise ValueError("图片尺寸无效，无法生成缩略图")
            image.save(staged_path, format="WEBP", quality=80, method=6)
        folder.mkdir(parents=True, exist_ok=True)
        os.replace(staged_path, final_path)
        return str(final_path), width, height
    finally:
        staged_path.unlink(missing_ok=True)
        try:
            staging_dir.rmdir()
        except OSError:
            pass


def assert_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("只允许下载公开的 http/https 图片")
    for result in socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM):
        address = ipaddress.ip_address(result[4][0])
        if address.is_private or address.is_loopback or address.is_link_local or address.is_reserved or address.is_multicast:
            raise ValueError("拒绝访问本机或内网地址")


class PublicRedirectHandler(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, request, fp, code, message, headers, new_url):
        assert_public_url(new_url)
        return super().redirect_request(request, fp, code, message, headers, new_url)


def download_image(material: dict) -> tuple[bytes, str]:
    media_url = material.get("thumbnail_url") or material.get("source_url")
    assert_public_url(media_url)
    request = urllib.request.Request(media_url, headers={"User-Agent": "MuseVault/0.2 single-user-capture"})
    opener = urllib.request.build_opener(PublicRedirectHandler())
    with opener.open(request, timeout=30) as response:
        assert_public_url(response.geturl())
        content_type = response.headers.get_content_type()
        if not content_type.startswith("image/"):
            raise ValueError(f"远程内容不是图片：{content_type}")
        chunks, size = [], 0
        while True:
            chunk = response.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_IMAGE_BYTES:
                raise ValueError("图片超过 30MB 限制")
            chunks.append(chunk)
    return b"".join(chunks), content_type


def export_readable_files(material: dict, formats: list[str]) -> tuple[str | None, str | None]:
    from docx import Document
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    folder = ASSET_ROOT / "web" / safe_folder_name(material.get("folder_id") or "inbox")
    folder.mkdir(parents=True, exist_ok=True)
    pdf_path = folder / f"{material['id']}.pdf"
    docx_path = folder / f"{material['id']}.docx"
    text = (material.get("text_content") or "").strip()
    paragraphs = [part.strip() for part in re.split(r"\n{2,}", text) if part.strip()]
    if not paragraphs:
        paragraphs = [material.get("ai_summary") or "页面没有提取到可读正文。"]

    if "docx" in formats:
        document = Document()
        document.add_heading(material["title"], level=0)
        document.add_paragraph(f"来源：{material.get('source_url', '')}")
        if material.get("ai_summary"):
            document.add_heading("内容摘要", level=1)
            document.add_paragraph(material["ai_summary"])
        document.add_heading("网页正文", level=1)
        for paragraph in paragraphs[:500]:
            document.add_paragraph(paragraph[:12000])
        document.save(docx_path)

    if "pdf" in formats:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("ChineseTitle", parent=styles["Title"], fontName="STSong-Light", fontSize=18, leading=25, spaceAfter=12)
        body_style = ParagraphStyle("ChineseBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=10.5, leading=17, wordWrap="CJK", spaceAfter=7)
        story = [Paragraph(html.escape(material["title"]), title_style), Paragraph(html.escape(f"来源：{material.get('source_url', '')}"), body_style), Spacer(1, 4 * mm)]
        if material.get("ai_summary"):
            story.extend([Paragraph("内容摘要", title_style), Paragraph(html.escape(material["ai_summary"]), body_style)])
        story.append(Paragraph("网页正文", title_style))
        story.extend(Paragraph(html.escape(paragraph[:12000]).replace("\n", "<br/>"), body_style) for paragraph in paragraphs[:500])
        SimpleDocTemplate(str(pdf_path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title=material["title"]).build(story)
    return str(pdf_path) if "pdf" in formats else None, str(docx_path) if "docx" in formats else None


def merge_tag_sources(user_tags: list[str], ai_tags: list[str]) -> list[str]:
    result: list[str] = []
    for value in [*user_tags, *ai_tags]:
        clean = " ".join(str(value).split()).strip()
        if clean and clean not in result:
            result.append(clean)
    return result[:20]


def finalize_material(conn: sqlite3.Connection, material_id: str) -> str:
    material = conn.execute("SELECT local_path, snapshot_path, source_url FROM materials WHERE id=?", (material_id,)).fetchone()
    if not material:
        return "missing"
    dead = conn.execute("SELECT stage, last_error FROM jobs WHERE material_id=? AND state='dead'", (material_id,)).fetchall()
    pending = conn.execute("SELECT COUNT(*) FROM jobs WHERE material_id=? AND state IN ('queued','running')", (material_id,)).fetchone()[0]
    if dead:
        status = "partial" if (material["local_path"] or material["snapshot_path"] or material["source_url"]) else "failed"
        reason = "必需任务失败：" + ", ".join(item["stage"] for item in dead)
    elif pending:
        status, reason = "analyzing", None
    else:
        status, reason = "ready", None
    conn.execute("UPDATE materials SET status=?, status_reason=? WHERE id=?", (status, reason, material_id))
    return status


def claim_download_operation(material: dict) -> dict:
    operation_id = str(uuid.uuid4())
    staged = STAGING_ROOT / operation_id / "local.part"
    with db() as conn:
        conn.execute(
            """INSERT INTO capture_operations
               (id, material_id, state, local_staged_path, created_at, updated_at)
               VALUES (?, ?, 'pending', ?, ?, ?)""",
            (operation_id, material["id"], str(staged), now(), now()),
        )
    return {"id": operation_id, "staged": staged}


def commit_download_operation(operation: dict, material: dict, content: bytes, mime: str) -> None:
    extension = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp", "image/gif": "gif", "image/avif": "avif"}.get(mime, "img")
    final = ASSET_ROOT / "image" / safe_folder_name(material.get("folder_id") or "inbox") / f"{material['id']}.{extension}"
    write_staged_file(operation["staged"], content)
    final.parent.mkdir(parents=True, exist_ok=True)
    with db() as conn:
        conn.execute("UPDATE capture_operations SET local_final_path=?, updated_at=? WHERE id=?", (str(final), now(), operation["id"]))
    os.replace(operation["staged"], final)
    with db() as conn:
        conn.execute("UPDATE capture_operations SET state='committed', updated_at=? WHERE id=?", (now(), operation["id"]))
        conn.execute("UPDATE materials SET local_path=?, content_hash=?, status='queued', status_reason=NULL WHERE id=?", (str(final), sha256_bytes(content), material["id"]))
        conn.execute("INSERT OR IGNORE INTO material_derivatives (material_id) VALUES (?)", (material["id"],))
        conn.execute("UPDATE material_derivatives SET state='queued', error=NULL, updated_at=? WHERE material_id=?", (now(), material["id"]))
        queue_job(conn, material["id"], "thumbnail")
        queue_job(conn, material["id"], "visual")
        conn.execute("UPDATE jobs SET state='done', last_error=NULL WHERE material_id=? AND stage='download'", (material["id"],))
        finalize_material(conn, material["id"])


def process_long_stage(job: sqlite3.Row, material: dict) -> None:
    """Run image/AI/export work without holding a SQLite transaction."""
    stage = job["stage"]

    if stage == "thumbnail":
        with db() as conn:
            conn.execute(
                "UPDATE material_derivatives SET state='running', error=NULL, updated_at=? WHERE material_id=?",
                (now(), material["id"]),
            )
        thumbnail_path, width, height = generate_thumbnail(material)
        with db() as conn:
            conn.execute(
                """UPDATE material_derivatives
                   SET thumbnail_path=?, thumbnail_width=?, thumbnail_height=?, state='ready', error=NULL, updated_at=?
                   WHERE material_id=?""",
                (thumbnail_path, width, height, now(), material["id"]),
            )
        return

    if stage == "visual":
        if not material["local_path"] or not Path(material["local_path"]).exists():
            raise FileNotFoundError("visual analysis source is missing")
        palette, histogram, brightness, detail = analyze_visual(material["local_path"])
        with db() as conn:
            conn.execute(
                "UPDATE visual_features SET palette_json=?, histogram_json=?, mean_brightness=?, detail_score=? WHERE material_id=?",
                (json.dumps(palette), json.dumps(histogram), brightness, detail, material["id"]),
            )
        return

    if stage == "ai_enrich":
        with db() as conn:
            row = conn.execute("SELECT * FROM materials WHERE id=?", (material["id"],)).fetchone()
        if not row:
            raise FileNotFoundError("AI analysis material is missing")
        fresh = dict(row)
        result = analyze_material(fresh)
        usage = result.pop("_usage", {})
        result.pop("_model", None)
        with db() as conn:
            latest_row = conn.execute("SELECT * FROM materials WHERE id=?", (material["id"],)).fetchone()
            if not latest_row:
                raise FileNotFoundError("AI analysis material was removed")
            latest = dict(latest_row)
            user_tags = json.loads(latest.get("tags_user_json") or "[]")
            rejected = set(json.loads(latest.get("tags_rejected_json") or "[]"))
            ai_tags = [tag for tag in compact_tags(result.get("tags", [])) if tag not in rejected]
            category = result.get("category") if result.get("category") in {"鍥剧墖", "瑙嗛", "鏂囨。", "缃戦〉"} else latest["category"]
            conn.execute(
                """UPDATE materials SET ai_summary=?, tags_json=?, entities_json=?, category=?, subcategory=?, status='analyzing',
                                tags_ai_json=? WHERE id=?""",
                (
                    str(result.get("summary") or latest["ai_summary"] or "")[:2000],
                    json.dumps(merge_tag_sources(user_tags, ai_tags), ensure_ascii=False),
                    json.dumps(result.get("entities", []), ensure_ascii=False),
                    category,
                    str(result.get("subcategory") or latest["subcategory"] or "")[:100] or None,
                    json.dumps(ai_tags, ensure_ascii=False),
                    material["id"],
                ),
            )
            conn.execute(
                "UPDATE visual_features SET visual_description=? WHERE material_id=?",
                (image_prompt_bundle(result) if material.get("material_type") == "image" else str(result.get("visual_description") or "")[:2000], material["id"]),
            )
            finish_ai_batch_item(conn, job["batch_id"], material["id"], state="done", usage=usage)
        return

    if stage == "export":
        with db() as conn:
            conn.execute(
                "UPDATE offline_exports SET state='running', error=NULL, updated_at=? WHERE material_id=?",
                (now(), material["id"]),
            )
            row = conn.execute("SELECT * FROM materials WHERE id=?", (material["id"],)).fetchone()
            export_row = conn.execute(
                "SELECT requested_formats_json, pdf_path, docx_path FROM offline_exports WHERE material_id=?",
                (material["id"],),
            ).fetchone()
        if not row:
            raise FileNotFoundError("export material is missing")
        fresh = dict(row)
        requested = json.loads(export_row["requested_formats_json"] or "[]") if export_row else []
        formats = [
            value for value in requested
            if value in {"pdf", "docx"} and not export_row[f"{value}_path"]
        ] if export_row else []
        if formats:
            pdf_path, docx_path = export_readable_files(fresh, formats)
        else:
            pdf_path, docx_path = None, None
        with db() as conn:
            conn.execute(
                """UPDATE offline_exports SET pdf_path=COALESCE(?, pdf_path), docx_path=COALESCE(?, docx_path),
                                state='ready', error=NULL, updated_at=? WHERE material_id=?""",
                (pdf_path, docx_path, now(), material["id"]),
            )
        return

    raise ValueError(f"unsupported long stage: {stage}")


def process_one() -> bool:
    global _restart_recovery_done
    if not _restart_recovery_done:
        recover_ai_batches_after_restart()
    schedule_auto_batch()
    with db() as conn:
        job = conn.execute("""SELECT jobs.* FROM jobs
                               JOIN materials active ON active.id=jobs.material_id
                               WHERE jobs.state='queued' AND jobs.next_run_at <= ? AND active.deleted_at IS NULL
                               AND (jobs.stage <> 'ai_enrich' OR jobs.batch_id IS NULL OR EXISTS (
                                 SELECT 1 FROM ai_batch_runs batch_guard
                                 WHERE batch_guard.id=jobs.batch_id AND batch_guard.state IN ('queued','running')
                               ))
                               AND (stage <> 'index' OR NOT EXISTS (
                                 SELECT 1 FROM jobs sibling WHERE sibling.material_id=jobs.material_id
                                 AND sibling.id<>jobs.id AND sibling.state IN ('queued','running')
                               ))
                               ORDER BY CASE stage
                                WHEN 'download' THEN 0 WHEN 'classify' THEN 1 WHEN 'thumbnail' THEN 2 WHEN 'visual' THEN 3
                                WHEN 'ai_enrich' THEN 4 WHEN 'export' THEN 5 WHEN 'transcribe' THEN 6 ELSE 7 END,
                                next_run_at LIMIT 1""", (now(),)).fetchone()
        if not job:
            return False
        conn.execute("UPDATE jobs SET state='running', attempts=attempts+1 WHERE id=?", (job["id"],))
        if job["stage"] == "ai_enrich":
            settings = get_processing_settings(conn)
            if ai_usage_today(conn)["requests"] >= settings["daily_request_limit"]:
                conn.execute("UPDATE jobs SET state='paused', last_error='已达到今日 AI 请求上限' WHERE id=?", (job["id"],))
                if job["batch_id"]:
                    conn.execute(
                        "UPDATE ai_batch_items SET state='paused', error='已达到今日 AI 请求上限' WHERE batch_id=? AND material_id=?",
                        (job["batch_id"], job["material_id"]),
                    )
                    conn.execute("UPDATE ai_batch_runs SET state='paused', stopped_at=? WHERE id=?", (now(), job["batch_id"]))
                return True
            start_ai_batch_item(conn, job["batch_id"], job["material_id"])
        material_row = conn.execute("SELECT * FROM materials WHERE id=?", (job["material_id"],)).fetchone()
    if not material_row:
        with db() as conn:
            conn.execute("UPDATE jobs SET state='dead', last_error='素材不存在' WHERE id=?", (job["id"],))
        return True
    material = dict(material_row)
    if material.get("deleted_at") is not None:
        with db() as conn:
            conn.execute("UPDATE jobs SET state='paused', last_error=NULL WHERE id=?", (job["id"],))
        return True

    try:
        if job["stage"] == "download":
            operation = claim_download_operation(material)
            try:
                content, mime = download_image(material)
                commit_download_operation(operation, material, content, mime)
                return True
            except Exception as error:
                mark_operation_failed(operation["id"], material["id"], error)
                with db() as conn:
                    attempts = job["attempts"] + 1
                    state = "dead" if attempts >= 3 else "queued"
                    conn.execute("UPDATE jobs SET state=?, next_run_at=?, last_error=? WHERE id=?", (state, now() + min(2 ** attempts * 30, 3600), redact_sensitive_text(error)[:500], job["id"]))
                    finalize_material(conn, material["id"])
                return True
        if job["stage"] in {"thumbnail", "visual", "ai_enrich", "export"}:
            process_long_stage(job, material)
            with db() as conn:
                conn.execute("UPDATE jobs SET state='done', last_error=NULL WHERE id=?", (job["id"],))
                finalize_material(conn, material["id"])
            return True
        with db() as conn:
            if job["stage"] == "classify":
                category, subcategory, tags = category_for("\n".join([material["title"], material["text_content"] or ""]), material["material_type"])
                summary = (material["text_content"] or material["title"]).strip().replace("\n", " ")[:320]
                user_tags = json.loads(material.get("tags_user_json") or "[]")
                ai_tags = compact_tags([*json.loads(material.get("tags_ai_json") or "[]"), *tags])
                rejected = set(json.loads(material.get("tags_rejected_json") or "[]"))
                ai_tags = [tag for tag in ai_tags if tag not in rejected]
                conn.execute("""UPDATE materials SET category=?, subcategory=?, tags_json=?, tags_ai_json=?, ai_summary=?, status='analyzing' WHERE id=?""",
                             (category, subcategory, json.dumps(merge_tag_sources(user_tags, ai_tags), ensure_ascii=False), json.dumps(ai_tags, ensure_ascii=False), summary, material["id"]))
            elif job["stage"] == "thumbnail":
                conn.execute(
                    "UPDATE material_derivatives SET state='running', error=NULL, updated_at=? WHERE material_id=?",
                    (now(), material["id"]),
                )
                thumbnail_path, width, height = generate_thumbnail(material)
                conn.execute(
                    """UPDATE material_derivatives
                       SET thumbnail_path=?, thumbnail_width=?, thumbnail_height=?, state='ready', error=NULL, updated_at=?
                       WHERE material_id=?""",
                    (thumbnail_path, width, height, now(), material["id"]),
                )
            elif job["stage"] == "visual":
                if not material["local_path"] or not Path(material["local_path"]).exists():
                    raise FileNotFoundError("视觉分析所需的本地预览文件不存在")
                palette, histogram, brightness, detail = analyze_visual(material["local_path"])
                conn.execute("UPDATE visual_features SET palette_json=?, histogram_json=?, mean_brightness=?, detail_score=? WHERE material_id=?",
                             (json.dumps(palette), json.dumps(histogram), brightness, detail, material["id"]))
            elif job["stage"] == "ai_enrich":
                fresh = dict(conn.execute("SELECT * FROM materials WHERE id=?", (material["id"],)).fetchone())
                result = analyze_material(fresh)
                usage = result.pop("_usage", {})
                result.pop("_model", None)
                user_tags = json.loads(fresh.get("tags_user_json") or "[]")
                rejected = set(json.loads(fresh.get("tags_rejected_json") or "[]"))
                ai_tags = [tag for tag in compact_tags(result.get("tags", [])) if tag not in rejected]
                category = result.get("category") if result.get("category") in {"图片", "视频", "文档", "网页"} else fresh["category"]
                conn.execute("""UPDATE materials SET ai_summary=?, tags_json=?, entities_json=?, category=?, subcategory=?, status='analyzing'
                                , tags_ai_json=?
                                WHERE id=?""",
                             (str(result.get("summary") or fresh["ai_summary"] or "")[:2000],
                              json.dumps(merge_tag_sources(user_tags, ai_tags), ensure_ascii=False),
                              json.dumps(result.get("entities", []), ensure_ascii=False), category,
                              str(result.get("subcategory") or fresh["subcategory"] or "")[:100] or None,
                              json.dumps(ai_tags, ensure_ascii=False), material["id"]))
                conn.execute("UPDATE visual_features SET visual_description=? WHERE material_id=?",
                             (image_prompt_bundle(result) if material.get("material_type") == "image" else str(result.get("visual_description") or "")[:2000], material["id"]))
                finish_ai_batch_item(
                    conn, job["batch_id"], material["id"], state="done", usage=usage,
                )
            elif job["stage"] == "export":
                conn.execute("UPDATE offline_exports SET state='running', error=NULL, updated_at=? WHERE material_id=?", (now(), material["id"]))
                fresh = dict(conn.execute("SELECT * FROM materials WHERE id=?", (material["id"],)).fetchone())
                export_row = conn.execute(
                    "SELECT requested_formats_json, pdf_path, docx_path FROM offline_exports WHERE material_id=?",
                    (material["id"],),
                ).fetchone()
                requested = json.loads(export_row["requested_formats_json"] or "[]") if export_row else []
                formats = [
                    value for value in requested
                    if value in {"pdf", "docx"} and not export_row[f"{value}_path"]
                ] if export_row else []
                if not formats:
                    conn.execute("UPDATE offline_exports SET state='ready', error=NULL, updated_at=? WHERE material_id=?", (now(), material["id"]))
                    formats = []
                pdf_path, docx_path = export_readable_files(fresh, formats)
                conn.execute("""UPDATE offline_exports SET pdf_path=COALESCE(?, pdf_path), docx_path=COALESCE(?, docx_path),
                                state='ready', error=NULL, updated_at=? WHERE material_id=?""",
                             (pdf_path, docx_path, now(), material["id"]))
            elif job["stage"] == "index":
                update_fts(conn, material["id"])
            elif job["stage"] == "transcribe":
                conn.execute("UPDATE transcripts SET state='waiting_for_local_asr' WHERE material_id=?", (material["id"],))
            conn.execute("UPDATE jobs SET state='done', last_error=NULL WHERE id=?", (job["id"],))
            finalize_material(conn, material["id"])
        return True
    except Exception as exc:
        if job["stage"] == "ai_enrich":
            # A provider/model/quota error is not a normal retryable worker
            # failure. Stop the batch on the first error so the user can fix
            # the model or quota before another request is billed.
            message = redact_sensitive_text(exc)[:500]
            with db() as conn:
                conn.execute(
                    "UPDATE jobs SET state='dead', last_error=?, next_run_at=? WHERE id=?",
                    (message, now(), job["id"]),
                )
                if job["batch_id"]:
                    conn.execute(
                        "UPDATE jobs SET state='paused', last_error=? WHERE batch_id=? AND stage='ai_enrich' AND state IN ('queued','running') AND id<>?",
                        ("因 AI 请求失败而暂停，修复配置后重试", job["batch_id"], job["id"]),
                    )
                    conn.execute(
                        "UPDATE ai_batch_items SET state='paused', error=? WHERE batch_id=? AND state IN ('queued','running') AND material_id<>?",
                        ("因 AI 请求失败而暂停，修复配置后重试", job["batch_id"], material["id"]),
                    )
                    finish_ai_batch_item(
                        conn, job["batch_id"], material["id"],
                        state="failed", error=message,
                    )
                    conn.execute(
                        "UPDATE ai_batch_runs SET state='paused', stopped_at=? WHERE id=?",
                        (now(), job["batch_id"]),
                    )
                finalize_material(conn, material["id"])
            return True
        with db() as conn:
            attempts = job["attempts"] + 1
            state = "dead" if attempts >= 3 else "queued"
            message = redact_sensitive_text(exc)[:500]
            conn.execute("UPDATE jobs SET state=?, next_run_at=?, last_error=? WHERE id=?",
                         (state, now() + min(2 ** attempts * 30, 3600), message, job["id"]))
            if job["stage"] == "export":
                conn.execute("UPDATE offline_exports SET state=?, error=?, updated_at=? WHERE material_id=?",
                             ("error" if state == "dead" else "queued", message, now(), material["id"]))
            if job["stage"] == "thumbnail":
                conn.execute(
                    "UPDATE material_derivatives SET state=?, error=?, updated_at=? WHERE material_id=?",
                    ("error" if state == "dead" else "queued", message, now(), material["id"]),
                )
            if job["stage"] == "ai_enrich":
                finish_ai_batch_item(
                    conn, job["batch_id"], material["id"],
                    state="failed" if state == "dead" else "queued", error=message,
                )
            finalize_material(conn, material["id"])
        return True


def recover_ai_batches_after_restart() -> None:
    """Require explicit user action after a worker restart before billing AI again."""
    global _restart_recovery_done
    if _restart_recovery_done:
        return
    with db() as conn:
        batches = conn.execute(
            "SELECT id FROM ai_batch_runs WHERE state IN ('queued','running','stopping')"
        ).fetchall()
        if not batches:
            _restart_recovery_done = True
            return
        for batch in batches:
            batch_id = batch["id"]
            reason = "程序重启后批次已暂停，请确认模型与额度后手动恢复"
            conn.execute(
                """UPDATE jobs SET state='paused', last_error=?
                   WHERE batch_id=? AND stage='ai_enrich' AND state IN ('queued','running')""",
                (reason, batch_id),
            )
            conn.execute(
                """UPDATE ai_batch_items SET state='paused', error=?
                   WHERE batch_id=? AND state IN ('queued','running')""",
                (reason, batch_id),
            )
            conn.execute(
                "UPDATE ai_batch_runs SET state='paused', stopped_at=? WHERE id=?",
                (now(), batch_id),
            )
            refresh_ai_batch(conn, batch_id)
    _restart_recovery_done = True


if __name__ == "__main__":
    init_db(repair=False)
    recover_ai_batches_after_restart()
    while True:
        if not process_one():
            time.sleep(1)
